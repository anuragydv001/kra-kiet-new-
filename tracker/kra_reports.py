"""
tracker/kra_reports.py
Unified report generation module for all KRAs.
Generates DOCX and PDF exports based on structured form data submitted by Deans.
"""

import os
import io
from datetime import date
from django.conf import settings

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

KIET_BLUE = RGBColor(0, 35, 111)
KIET_MAROON = RGBColor(98, 36, 35)
KIET_LIGHT = RGBColor(232, 237, 247)

def _get_logo_path() -> str:
    possible = [
        os.path.join(settings.BASE_DIR, 'frontend', 'public', 'images', 'kiet_university_logo.png'),
        r"d:\KIET-KRA--main\frontend\public\images\kiet_university_logo.png",
    ]
    for p in possible:
        if os.path.exists(p):
            return p
    return ""

def _get_logo_b64() -> str:
    possible = [
        os.path.join(settings.BASE_DIR, 'frontend', 'public', 'images', 'kiet_logo_b64.txt'),
        r"d:\KIET-KRA--main\frontend\public\images\kiet_logo_b64.txt",
    ]
    for p in possible:
        if os.path.exists(p):
            try:
                with open(p, 'r', encoding='utf-8') as f:
                    return f.read().strip()
            except Exception:
                pass
    return ""

def _normalize_rows(raw_rows, expected_keys, expected_cols_count):
    """
    Normalizes rows into a list of lists of strings, handling both list of dicts
    and list of lists/tuples, ensuring S.No. is at index 0.
    """
    rows = []
    for i, item in enumerate(raw_rows):
        if isinstance(item, dict):
            row = [str(i + 1)] + [str(item.get(k, '') if item.get(k) is not None else '') for k in expected_keys]
            rows.append(row)
        elif isinstance(item, (list, tuple)):
            if len(item) >= expected_cols_count:
                rows.append([str(x) if x is not None else '' for x in item[:expected_cols_count]])
            elif len(item) == expected_cols_count - 1:
                rows.append([str(i + 1)] + [str(x) if x is not None else '' for x in item])
            else:
                rows.append([str(x) if x is not None else '' for x in item])
    return rows

def _set_cell_bg(cell, hex_color: str):
    """Set background color for a Word table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _make_paragraph(para, text: str, bold=False, size_pt=10,
                    align=WD_ALIGN_PARAGRAPH.LEFT, color: RGBColor = None, italic=False):
    para.clear()
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    para.alignment = align
    return para


def _add_header(doc: Document, dept: str, kra_num: str, title: str,
                prepared_by: str, designation: str, cycle: str, academic_year: str, report_date: str):
    """Add standard KIET header and info block to a document."""
    logo_path = _get_logo_path()
    if logo_path:
        logo_para = doc.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_para.paragraph_format.space_before = Pt(0)
        logo_para.paragraph_format.space_after = Pt(2)
        run = logo_para.add_run()
        run.add_picture(logo_path, width=Cm(6.0))

    # Institution heading
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(2)
    heading.paragraph_format.space_after = Pt(1)
    _make_paragraph(heading, 'Krishna Institute of Engineering & Technology (KIET) (Deemed to be university)',
                    bold=True, size_pt=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0, 0, 0))
    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(1)
    sub.paragraph_format.space_after = Pt(4)
    _make_paragraph(sub, 'Delhi-NCR, Ghaziabad — ISO 9001:2015 Certified', size_pt=9,
                    align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(80, 80, 80))

    # Divider
    border_para = doc.add_paragraph()
    border_para.paragraph_format.space_after = Pt(4)
    pPr = border_para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '16')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), '622423')
    pBdr.append(bottom)
    pPr.append(pBdr)

    # Report title
    t_para = doc.add_paragraph()
    _make_paragraph(t_para, title.upper(), bold=True, size_pt=12,
                    align=WD_ALIGN_PARAGRAPH.CENTER)
    kn_para = doc.add_paragraph()
    _make_paragraph(kn_para, f'KRA {kra_num}', bold=True, size_pt=11,
                    align=WD_ALIGN_PARAGRAPH.CENTER, color=KIET_MAROON)

    # Info table
    tbl = doc.add_table(rows=3, cols=4)
    tbl.style = 'Table Grid'
    data = [
        ['Name of School/Department', dept, 'Academic Session', academic_year],
        ['Reporting Period (Bi-Monthly Cycle)', cycle, 'Date of Report', report_date],
        ['Report Prepared By', prepared_by, 'Designation', designation],
    ]
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx]
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = cell_text
            if c_idx in (0, 2):
                cell.paragraphs[0].runs[0].bold = True
                cell.paragraphs[0].runs[0].font.size = Pt(9)
                _set_cell_bg(cell, 'F7F7F7')
            else:
                cell.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph()


def _add_section_heading(doc: Document, text: str):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = KIET_MAROON
    shading = para._p.get_or_add_pPr()
    # inline highlight using character shading isn't trivial; use bold + color only


def _add_table(doc: Document, columns: list, rows: list):
    """Generic bordered table with header row."""
    if not columns:
        return
    tbl = doc.add_table(rows=1 + len(rows), cols=len(columns))
    tbl.style = 'Table Grid'
    # Header
    hdr = tbl.rows[0]
    for c_idx, col in enumerate(columns):
        cell = hdr.cells[c_idx]
        cell.text = col
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        _set_cell_bg(cell, 'E8EDF7')
    # Rows
    for r_idx, row_data in enumerate(rows):
        row = tbl.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val) if val is not None else ''
            cell.paragraphs[0].runs[0].font.size = Pt(8)
    doc.add_paragraph()


def _add_kv_table(doc: Document, items: list):
    """Key-value table for extra sections."""
    tbl = doc.add_table(rows=len(items), cols=2)
    tbl.style = 'Table Grid'
    for r_idx, (label, value) in enumerate(items):
        row = tbl.rows[r_idx]
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(9)
        _set_cell_bg(row.cells[0], 'E8EDF7')
        row.cells[1].text = str(value) if value is not None else ''
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(9)
        # set column widths
        row.cells[0].width = Cm(7)
    doc.add_paragraph()


def _add_signatures(doc: Document, prepared_by: str, designation: str, dept: str, report_date: str):
    """Two-column signature block."""
    tbl = doc.add_table(rows=1, cols=2)
    left = tbl.rows[0].cells[0]
    right = tbl.rows[0].cells[1]
    for cell, name, role in [(left, prepared_by, designation), (right, 'HOD / Dean', dept)]:
        p = cell.add_paragraph()
        p.add_run('Signature: ' + '_' * 25 + '\n\n').font.size = Pt(9)
        r = p.add_run(f'Name: {name}\nDesignation: {role}\nDate: {report_date}')
        r.font.size = Pt(9)


# ─────────────────────────────────────────────────────────────────────────────
# Context builders — produce the template context for the generic HTML template
# ─────────────────────────────────────────────────────────────────────────────

def _base_ctx(data: dict, kra_num: str, title: str) -> dict:
    today = date.today().strftime('%d/%m/%Y')
    return {
        'kra_number': kra_num,
        'report_title': title,
        'department': data.get('department', ''),
        'academic_year': data.get('academicYear', data.get('academic_year', '2026-2027')),
        'cycle': data.get('cycle', ''),
        'report_date': data.get('reportDate', today),
        'prepared_by': data.get('preparedBy', data.get('facultyName', data.get('coordinatorName', ''))),
        'designation': data.get('designation', 'Professor & Dean'),
    }


def build_generic_ctx(data: dict, kra_num: str, title: str) -> dict:
    """
    Build a template context for the generic HTML PDF template.
    data must contain: department, cycle, academicYear, columns, rows, 
    and optionally: tablTitle, extraSections, narrative, surveyQuestions,
    overallSatisfaction, suggestions, ratingOptions.
    """
    ctx = _base_ctx(data, kra_num, title)
    ctx['table_title'] = data.get('tableTitle', 'Activity Details')
    ctx['columns'] = data.get('columns', [])
    ctx['rows'] = data.get('rows', [])
    ctx['survey_questions'] = data.get('surveyQuestions', [])
    ctx['rating_options'] = data.get('ratingOptions', ['Poor(1)', 'Satisfactory(2)', 'Good(3)', 'Very Good(4)', 'Excellent(5)'])
    ctx['overall_satisfaction'] = data.get('overallSatisfaction', '')
    ctx['suggestions'] = data.get('suggestions', '')
    ctx['extra_sections'] = data.get('extraSections', [])
    ctx['narrative'] = data.get('narrative', '')
    ctx['logo_b64'] = _get_logo_b64()
    return ctx


# ─────────────────────────────────────────────────────────────────────────────
# DOCX generators — one per KRA family
# ─────────────────────────────────────────────────────────────────────────────

KRA_TITLES = {
    '8': 'Curriculum Feedback – Employer and Industry',
    '9': 'Curriculum Feedback – Alumni',
    '10': 'Curriculum Feedback – Academic Expert (Outside)',
    '11': 'Report on Bridge Courses Conducted',
    '13': 'Report on Industry Advisory Board',
    '18': 'Activity Report – Functional MoU with Industry / Research Organization',
    '19': 'Industrial Visit Report',
    '23': 'Report on Skill Development and Employability Training',
    '33': 'Alumni Events Organized',
    '36': 'Report on Administrative Training Programs (FDP / STP)',
    '37': 'Report on Faculty Development Programs',
    '44': 'Report on BOOT Camps and Peer Learning Activities',
    '45': 'National and International Flagship Event Report',
}


def _feedback_docx(data: dict, kra_num: str, title: str, respondent_label: str) -> bytes:
    """
    Generic DOCX for structured feedback KRAs (8, 9, 10).
    data keys: department, cycle, academicYear, respondentName, designation,
               email, mobileNumber, organization, subject,
               ratings (dict {q_id: rating_value}),
               overallSatisfaction, suggestions, surveyQuestions
    """
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), kra_num, title,
                data.get('respondentName', ''), data.get('designation', respondent_label),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'),
                data.get('reportDate', today))

    _add_section_heading(doc, f'{respondent_label} Details')
    _add_kv_table(doc, [
        ('Name', data.get('respondentName', '')),
        ('Designation / Role', data.get('designation', respondent_label)),
        ('Organization / Company', data.get('organization', '')),
        ('Email', data.get('email', '')),
        ('Mobile No.', data.get('mobileNumber', '')),
        ('Subject / Programme Reviewed', data.get('subject', '')),
    ])

    questions = data.get('surveyQuestions', [])
    ratings = data.get('ratings', {})
    rating_labels = {1: 'Poor', 2: 'Satisfactory', 3: 'Good', 4: 'Very Good', 5: 'Excellent'}

    if questions:
        _add_section_heading(doc, 'Feedback Survey Responses')
        cols = ['S.No.', 'Survey Statement', 'Rating (1–5)', 'Rating Label']
        rows = []
        for q in questions:
            qid = str(q.get('id', ''))
            rating_val = ratings.get(qid) or ratings.get(int(qid) if qid.isdigit() else qid)
            rows.append([
                q.get('id', ''),
                q.get('statement', ''),
                str(rating_val) if rating_val is not None else '—',
                rating_labels.get(rating_val, '—'),
            ])
        _add_table(doc, cols, rows)

        _add_kv_table(doc, [
            ('Overall Satisfaction (1–10)', str(data.get('overallSatisfaction', ''))),
            ('Suggestions / Comments', data.get('suggestions', '')),
        ])

    _add_signatures(doc, data.get('respondentName', ''), respondent_label,
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra8_docx(data: dict) -> bytes:
    return _feedback_docx(data, '8', KRA_TITLES['8'], 'Employer / Industry Representative')


def generate_kra9_docx(data: dict) -> bytes:
    suffix = data.get('alumniType', 'lt4')
    label = 'Alumni (< 4 Years)' if suffix == 'lt4' else 'Alumni (> 4 Years)'
    return _feedback_docx(data, '9', KRA_TITLES['9'], label)


def generate_kra10_docx(data: dict) -> bytes:
    return _feedback_docx(data, '10', KRA_TITLES['10'], 'Academic Expert (External)')


def generate_kra11_docx(data: dict) -> bytes:
    """KRA 11 — Bridge Courses."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), '11', KRA_TITLES['11'],
                data.get('coordinatorName', ''), data.get('designation', 'Faculty Coordinator'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'Details of Bridge Courses')
    courses = data.get('courses') or data.get('rows') or []
    cols = ['S.No.', 'Course Name', 'Duration (Hours)', 'Mode', 'Resource Person', 'Enrolled', 'Completed']
    keys = ['courseName', 'duration', 'mode', 'resourcePerson', 'enrolled', 'completed']
    rows = _normalize_rows(courses, keys, 7)
    _add_table(doc, cols, rows)

    syllabus = data.get('syllabus', [])
    if syllabus:
        _add_section_heading(doc, 'Course Content / Syllabus')
        _add_table(doc, ['Module', 'Topic', 'Hours'],
                   [[s.get('module', ''), s.get('topic', ''), str(s.get('hours', ''))] for s in syllabus])

    _add_section_heading(doc, 'Assessment and Certification')
    _add_kv_table(doc, [
        ('Assessment Method', data.get('assessmentMethod', '')),
        ('Minimum Attendance Requirement', data.get('minAttendance', '')),
        ('Passing / Completion Criteria', data.get('passingCriteria', '')),
        ('Certificate Issuing Authority', data.get('certAuthority', '')),
        ('Number of Certificates Issued', str(data.get('numCertsIssued', ''))),
    ])

    _add_section_heading(doc, 'Outcome and Impact')
    _add_kv_table(doc, [
        ('Overall Feedback Score', f"{data.get('feedbackScore', '')} / 5"),
        ('Major Suggestions Received', data.get('majorSuggestions', '')),
        ('Impact on Students', data.get('impact', '')),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Faculty Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra13_docx(data: dict) -> bytes:
    """KRA 13 — Industry Advisory Board."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), '13', KRA_TITLES['13'],
                data.get('coordinatorName', ''), data.get('designation', 'Professor & Dean'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'Meeting / Activity Details')
    _add_kv_table(doc, [
        ('Date of Meeting', data.get('meetingDate', '')),
        ('Venue', data.get('venue', '')),
        ('Mode (Online / Offline)', data.get('mode', '')),
        ('Agenda of Meeting', data.get('agenda', '')),
    ])

    members = data.get('members') or data.get('rows') or []
    if members:
        _add_section_heading(doc, 'Members of Industry Advisory Board')
        cols = ['S.No.', 'Name', 'Designation', 'Organization', 'Domain']
        keys = ['name', 'designation', 'organization', 'domain']
        rows = _normalize_rows(members, keys, 5)
        _add_table(doc, cols, rows)

    _add_section_heading(doc, 'Decisions and Recommendations')
    _add_kv_table(doc, [
        ('Key Recommendations', data.get('recommendations', '')),
        ('Action Points', data.get('actionPoints', '')),
        ('Follow-up Scheduled', data.get('followUp', '')),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Faculty Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra18_docx(data: dict) -> bytes:
    """KRA 18 — Functional MoU Activity Report."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), '18', KRA_TITLES['18'],
                data.get('coordinatorName', ''), data.get('designation', 'Professor & Dean'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'MoU Details')
    _add_kv_table(doc, [
        ('Name of Industry / Research Organization', data.get('industryName', '')),
        ('Date of MoU Signing', data.get('mouDate', '')),
        ('Duration / Validity of MoU', data.get('mouDuration', '')),
        ('Key Objectives / Scope', data.get('objectives', '')),
    ])

    activities = data.get('activities') or data.get('rows') or []
    if activities:
        _add_section_heading(doc, 'Activities Conducted under MoU')
        cols = ['S.No.', 'Activity Description', 'Date', 'No. of Beneficiaries', 'Outcome']
        keys = ['description', 'date', 'beneficiaries', 'outcome']
        rows = _normalize_rows(activities, keys, 5)
        _add_table(doc, cols, rows)

    _add_section_heading(doc, 'Impact Summary')
    _add_kv_table(doc, [
        ('Benefits to Students', data.get('studentBenefits', '')),
        ('Benefits to Faculty', data.get('facultyBenefits', '')),
        ('Overall Assessment', data.get('overallAssessment', '')),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Faculty Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra19_docx(data: dict) -> bytes:
    """KRA 19 — Industrial Visit Report."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), '19', KRA_TITLES['19'],
                data.get('coordinatorName', ''), data.get('designation', 'Faculty Coordinator'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'Visit Details')
    _add_kv_table(doc, [
        ('Name of Industry / Organization Visited', data.get('industryName', '')),
        ('Address / Location', data.get('address', '')),
        ('Date(s) of Visit', data.get('visitDate', '')),
        ('Duration of Visit', data.get('duration', '')),
        ('Mode (On-Site / Virtual)', data.get('mode', 'On-Site')),
        ('Year / Semester of Students', data.get('studentYear', '')),
        ('Faculty Escort', data.get('facultyEscort', '')),
    ])

    students = data.get('students') or data.get('rows') or []
    if students:
        _add_section_heading(doc, 'Student Participants')
        cols = ['S.No.', 'Student Name', 'Roll No.', 'Year/Sem', 'Feedback Score (1-5)']
        keys = ['name', 'rollNo', 'year', 'feedbackScore']
        rows = _normalize_rows(students, keys, 5)
        _add_table(doc, cols, rows)

    _add_section_heading(doc, 'Learning Outcomes and Student Feedback')
    _add_kv_table(doc, [
        ('Key Learning Outcomes', data.get('learningOutcomes', '')),
        ('Student Feedback Summary', data.get('feedbackSummary', '')),
        ('Overall Feedback Score (avg)', str(data.get('overallScore', ''))),
        ('Impact on Curriculum / Placement', data.get('impact', '')),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Faculty Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra23_docx(data: dict) -> bytes:
    """KRA 23 — Skill Development & Employability Training."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), '23', KRA_TITLES['23'],
                data.get('coordinatorName', ''), data.get('designation', 'Faculty Coordinator'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'Training Program Details')
    _add_kv_table(doc, [
        ('Program Title', data.get('programTitle', '')),
        ('Organizing Body / Partner', data.get('organizingBody', '')),
        ('Date(s)', data.get('dates', '')),
        ('Duration (Hours)', str(data.get('durationHours', ''))),
        ('Mode (Offline / Online / Hybrid)', data.get('mode', '')),
        ('Target Audience', data.get('targetAudience', '')),
        ('No. of Students Enrolled', str(data.get('enrolled', ''))),
        ('No. of Students Completed', str(data.get('completed', ''))),
    ])

    sessions = data.get('sessions') or data.get('rows') or []
    if sessions:
        _add_section_heading(doc, 'Session-wise Details')
        cols = ['S.No.', 'Session Topic', 'Resource Person', 'Date', 'Duration (Hrs)']
        keys = ['topic', 'resourcePerson', 'date', 'hours']
        rows = _normalize_rows(sessions, keys, 5)
        _add_table(doc, cols, rows)

    _add_section_heading(doc, 'Outcome and Impact')
    _add_kv_table(doc, [
        ('Skills Imparted', data.get('skillsImparted', '')),
        ('Placement Impact', data.get('placementImpact', '')),
        ('Overall Feedback Score (avg, out of 5)', str(data.get('feedbackScore', ''))),
        ('Major Suggestions', data.get('suggestions', '')),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Faculty Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra33_docx(data: dict) -> bytes:
    """KRA 33 — Alumni Events Organized."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), '33', KRA_TITLES['33'],
                data.get('coordinatorName', ''), data.get('designation', 'Faculty Coordinator'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'Event Details')
    _add_kv_table(doc, [
        ('Event Title', data.get('eventTitle', '')),
        ('Date of Event', data.get('eventDate', '')),
        ('Venue', data.get('venue', '')),
        ('Mode (Online / Offline / Hybrid)', data.get('mode', '')),
        ('Number of Alumni Participated', str(data.get('alumniCount', ''))),
        ('Number of Students Participated', str(data.get('studentCount', ''))),
    ])

    alumni = data.get('alumni') or data.get('rows') or []
    if alumni:
        _add_section_heading(doc, 'List of Participating Alumni')
        cols = ['S.No.', 'Name', 'Batch/Year', 'Current Designation', 'Organization', 'Contact']
        keys = ['name', 'batch', 'designation', 'organization', 'contact']
        rows = _normalize_rows(alumni, keys, 6)
        _add_table(doc, cols, rows)

    _add_section_heading(doc, 'Outcome and Impact')
    _add_kv_table(doc, [
        ('Key Highlights / Activities During Event', data.get('highlights', '')),
        ('Mentorship / Career Guidance Provided', data.get('mentorship', '')),
        ('Alumni Contribution / Pledge', data.get('contribution', '')),
        ('Overall Feedback / Observations', data.get('feedback', '')),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Faculty Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra36_37_docx(data: dict) -> bytes:
    """KRA 36/37 — FDP / STP / Administrative Training Report."""
    kra_num = data.get('kraNum', '36-37')
    title = 'Report on Faculty Development Programs (FDP) and Staff Training Programs (STP)'
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), kra_num, title,
                data.get('coordinatorName', ''), data.get('designation', 'Professor & Dean'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'Program Details')
    _add_kv_table(doc, [
        ('Program Title', data.get('programTitle', '')),
        ('Organizing Institution / Body', data.get('organizingBody', '')),
        ('Type (FDP / STP / Workshop / Admin Training)', data.get('programType', '')),
        ('Date(s)', data.get('dates', '')),
        ('Duration', data.get('duration', '')),
        ('Mode (Online / Offline / Hybrid)', data.get('mode', '')),
        ('Venue', data.get('venue', '')),
        ('Resource Person(s)', data.get('resourcePersons', '')),
    ])

    participants = data.get('participants') or data.get('rows') or []
    if participants:
        _add_section_heading(doc, 'Participants')
        cols = ['S.No.', 'Name', 'Designation', 'Department', 'Certificate Received']
        keys = ['name', 'designation', 'dept', 'certificateReceived']
        rows = _normalize_rows(participants, keys, 5)
        _add_table(doc, cols, rows)

    _add_section_heading(doc, 'Topics Covered / Modules')
    topics = data.get('topics', [])
    if topics:
        _add_table(doc, ['S.No.', 'Topic / Module', 'Resource Person', 'Duration (Hrs)'],
                   [[str(i + 1), t.get('topic', ''), t.get('person', ''), str(t.get('hours', ''))]
                    for i, t in enumerate(topics)])

    _add_section_heading(doc, 'Outcome and Impact')
    _add_kv_table(doc, [
        ('Key Skills / Knowledge Gained', data.get('skillsGained', '')),
        ('Impact on Teaching / Administration', data.get('impact', '')),
        ('Overall Feedback Score (avg, out of 5)', str(data.get('feedbackScore', ''))),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Program Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra44_docx(data: dict) -> bytes:
    """KRA 44 — BOOT Camps and Peer Learning Activities."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), '44', KRA_TITLES['44'],
                data.get('coordinatorName', ''), data.get('designation', 'Faculty Coordinator'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'Event Details')
    _add_kv_table(doc, [
        ('Event Title', data.get('eventTitle', '')),
        ('Type (BOOT Camp / Peer Learning / Both)', data.get('eventType', '')),
        ('Date(s)', data.get('dates', '')),
        ('Venue / Platform', data.get('venue', '')),
        ('Mode (Online / Offline / Hybrid)', data.get('mode', '')),
        ('No. of Students Participated', str(data.get('studentCount', ''))),
        ('Resource Person(s) / Mentor(s)', data.get('resourcePersons', '')),
    ])

    sessions = data.get('sessions') or data.get('rows') or []
    if sessions:
        _add_section_heading(doc, 'Session-wise Details')
        cols = ['S.No.', 'Session Topic', 'Mode', 'Date', 'Participants', 'Resource Person']
        keys = ['topic', 'mode', 'date', 'participants', 'resourcePerson']
        rows = _normalize_rows(sessions, keys, 6)
        _add_table(doc, cols, rows)

    _add_section_heading(doc, 'Outcome and Impact')
    _add_kv_table(doc, [
        ('Skills Developed', data.get('skillsDeveloped', '')),
        ('Project / Product Outcomes', data.get('projectOutcomes', '')),
        ('Placement / Certification Impact', data.get('placementImpact', '')),
        ('Overall Feedback Score (avg, out of 5)', str(data.get('feedbackScore', ''))),
        ('Major Suggestions', data.get('suggestions', '')),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Faculty Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra45_docx(data: dict) -> bytes:
    """KRA 45 — National & International Flagship Event."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    _add_header(doc, data.get('department', ''), '45', KRA_TITLES['45'],
                data.get('coordinatorName', ''), data.get('designation', 'Event Coordinator'),
                data.get('cycle', ''), data.get('academicYear', '2026-2027'), data.get('reportDate', today))

    _add_section_heading(doc, 'Event Details')
    _add_kv_table(doc, [
        ('Event Name', data.get('eventName', '')),
        ('Type (National / International)', data.get('eventType', '')),
        ('Date(s)', data.get('dates', '')),
        ('Venue / Platform', data.get('venue', '')),
        ('Organizing Department(s)', data.get('organizingDepts', data.get('department', ''))),
        ('Chief Guest(s) / Keynote Speaker(s)', data.get('chiefGuests', '')),
        ('Total Participants', str(data.get('totalParticipants', ''))),
        ('External Participants', str(data.get('externalParticipants', ''))),
        ('Sponsorship / Funding Details', data.get('sponsorship', '')),
    ])

    events = data.get('subEvents') or data.get('rows') or []
    if events:
        _add_section_heading(doc, 'Sub-Events / Technical Sessions')
        cols = ['S.No.', 'Event/Session Title', 'Date', 'Participants', 'Winners/Outcome']
        keys = ['title', 'date', 'participants', 'outcome']
        rows = _normalize_rows(events, keys, 5)
        _add_table(doc, cols, rows)

    _add_section_heading(doc, 'Impact and Media Coverage')
    _add_kv_table(doc, [
        ('Key Highlights', data.get('highlights', '')),
        ('Media Coverage / Press', data.get('mediaCoverage', '')),
        ('Industry / Academic Linkages Established', data.get('linkages', '')),
        ('Follow-up Actions', data.get('followUp', '')),
    ])

    _add_signatures(doc, data.get('coordinatorName', ''), 'Event Coordinator',
                    data.get('department', ''), data.get('reportDate', today))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Dispatcher
# ─────────────────────────────────────────────────────────────────────────────

DOCX_GENERATORS = {
    '8': generate_kra8_docx,
    '9': generate_kra9_docx,
    '10': generate_kra10_docx,
    '11': generate_kra11_docx,
    '13': generate_kra13_docx,
    '18': generate_kra18_docx,
    '19': generate_kra19_docx,
    '23': generate_kra23_docx,
    '33': generate_kra33_docx,
    '36': generate_kra36_37_docx,
    '37': generate_kra36_37_docx,
    '44': generate_kra44_docx,
    '45': generate_kra45_docx,
}


def get_kra_title(kra_num: str) -> str:
    return KRA_TITLES.get(kra_num, f'KRA {kra_num} Report')


def generate_generic_kra_docx(kra_num: str, data: dict) -> bytes:
    """Generic Word (.docx) export for any KRA with tables, extra sections, and narrative."""
    doc = Document()
    today = date.today().strftime('%d/%m/%Y')
    title = data.get('reportTitle') or get_kra_title(kra_num)
    _add_header(
        doc,
        data.get('department', ''),
        kra_num,
        title,
        data.get('preparedBy', data.get('coordinatorName', 'Faculty Coordinator')),
        data.get('designation', 'Professor & Dean'),
        data.get('cycle', ''),
        data.get('academicYear', '2026-2027'),
        data.get('reportDate', today)
    )

    columns = data.get('columns', [])
    rows = data.get('rows', [])
    table_title = data.get('tableTitle', 'Activity & Performance Records')

    if columns and rows:
        _add_section_heading(doc, table_title)
        _add_table(doc, columns, rows)

    extra_sections = data.get('extraSections', [])
    for sec in extra_sections:
        _add_section_heading(doc, sec.get('title', 'Details'))
        items = [(item.get('label', ''), item.get('value', '')) for item in sec.get('items', [])]
        if items:
            _add_kv_table(doc, items)

    narrative = data.get('narrative', '')
    if narrative:
        _add_section_heading(doc, 'Outcome and Observations')
        p = doc.add_paragraph()
        run = p.add_run(narrative)
        run.font.size = Pt(9.5)

    _add_signatures(
        doc,
        data.get('preparedBy', data.get('coordinatorName', 'Faculty Coordinator')),
        data.get('designation', 'Professor & Dean'),
        data.get('department', ''),
        data.get('reportDate', today)
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_kra_docx(kra_num: str, data: dict) -> bytes:
    """Dispatch to the correct DOCX generator based on KRA number, or fallback to generic generator."""
    gen = DOCX_GENERATORS.get(kra_num)
    if gen:
        return gen(data)
    return generate_generic_kra_docx(kra_num, data)
